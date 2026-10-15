"""Job catalog upload, .xls parsing, and timesheet safety checks."""
import io
import os
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DB_PATH = Path(tempfile.mkdtemp(prefix='timesheet-test-')) / 'timesheet.db'
os.environ['DATABASE_URL'] = 'sqlite:///' + DB_PATH.as_posix()
os.environ['MASTER_JOB_FILE'] = str(ROOT / 'missing-master-job-list.xls')
os.environ['ADMIN_PASSWORD'] = 'admin'

sys.path.insert(0, str(ROOT))
import app as app_module  # noqa: E402


def make_xlsx(path, rows, title=None):
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Jobs'
    start = 1
    if title:
        sheet.cell(1, 1, title)
        start = 2
    headers = ['Job Number', 'Job Name', 'Client', 'Task', 'County', 'Date']
    for col, header in enumerate(headers, start=1):
        sheet.cell(start, col, header)
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            sheet.cell(start + 1 + offset, col, value)
    workbook.save(path)


def snapshot_timesheets():
    rows = []
    for entry in app_module.Timesheet.query.order_by(app_module.Timesheet.id).all():
        rows.append({
            'id': entry.id,
            'employee_name': entry.employee_name,
            'job_number': entry.job_number,
            'job_name': entry.job_name,
            'hours': entry.hours,
            'description': entry.description,
            'work_date': entry.work_date.isoformat(),
            'is_billed': entry.is_billed,
        })
    return rows


def main():
    assert app_module.normalize_database_url('postgres://u:p@db/timesheet') == 'postgresql+psycopg://u:p@db/timesheet'
    assert app_module.normalize_database_url('postgresql://u:p@db/timesheet') == 'postgresql+psycopg://u:p@db/timesheet'
    assert app_module.normalize_database_url('postgresql+psycopg://u:p@db/timesheet') == 'postgresql+psycopg://u:p@db/timesheet'
    assert app_module.normalize_database_url('sqlite:///timesheet.db') == 'sqlite:///timesheet.db'

    client = app_module.app.test_client()
    with app_module.app.app_context():
        before_seed = app_module.get_job_list()
        assert any(job['number'] == '126021' for job in before_seed), 'json seed missing'
        assert all(job['number'] != '126028' for job in before_seed)

        entry = app_module.Timesheet(
            employee_name='Brian Clark',
            job_number='122001',
            job_name='Forest Lake',
            hours=4.5,
            description='• Existing work that must stay',
            work_date=date(2026, 1, 15),
            is_billed=True,
        )
        app_module.db.session.add(entry)
        app_module.db.session.commit()
        entry_id = entry.id
        original_rows = snapshot_timesheets()

    login = client.post('/login', data={'password': 'admin'})
    assert login.status_code in (302, 303), login.status_code

    junk = client.post(
        '/admin/upload-jobs',
        data={'file': (io.BytesIO(b'this is not an excel file'), 'Job List - 09302009.xls')},
        content_type='multipart/form-data',
    )
    assert junk.status_code >= 400, junk.get_data(as_text=True)
    with app_module.app.app_context():
        assert snapshot_timesheets() == original_rows
        assert any(job['number'] == '126021' for job in app_module.get_job_list())

    xls_path = ROOT / 'tests' / 'fixtures' / 'job_list_09302009.xls'
    xlsx_path = DB_PATH.parent / 'jobs.xlsx'
    make_xlsx(
        xlsx_path,
        [(126028, '213 Hillandale', 'Hillandale LLC', 'Survey', 'Mecklenburg', '2009-09-30')],
        title='Southeast Geomatics',
    )

    with open(xls_path, 'rb') as handle:
        uploaded = client.post(
            '/admin/upload-jobs',
            data={'file': (handle, 'Job List - 09302009.xls')},
            content_type='multipart/form-data',
        )
    body = uploaded.get_json()
    assert uploaded.status_code == 200, body
    assert body['success'] is True
    assert body['count'] == 3

    with app_module.app.app_context():
        jobs = {job['number']: job['name'] for job in app_module.get_job_list()}
        assert jobs['126028'] == '213 Hillandale'
        assert jobs['H1000'] == 'Overhead'
        assert '126021' not in jobs
        assert snapshot_timesheets() == original_rows
        meta = app_module.db.session.get(app_module.JobCatalogMeta, 1)
        assert meta.source == 'upload'

    api = client.get('/api/jobs')
    assert 'no-store' in api.headers.get('Cache-Control', '')
    payload = {job['number']: job['name'] for job in api.get_json()}
    assert payload['126028'] == '213 Hillandale'

    page = client.get('/')
    html = page.get_data(as_text=True)
    assert '213 Hillandale' in html
    assert 'no-store' in page.headers.get('Cache-Control', '')

    admin = client.get('/admin')
    admin_html = admin.get_data(as_text=True)
    assert admin.status_code == 200
    assert 'Brian Clark' in admin_html
    assert 'size: letter' in admin_html
    assert 'Print</button>' in admin_html or '>Print<' in admin_html

    # An older master workbook must not replace the upload. A newer one still syncs.
    master_path = DB_PATH.parent / 'master.xlsx'
    make_xlsx(master_path, [(111111, 'Old Master Job', 'C', 'T', 'County', '2001-01-01')])
    os.utime(master_path, (1_000_000_000, 1_000_000_000))
    os.environ['MASTER_JOB_FILE'] = str(master_path)
    app_module._last_master_check = 0
    with app_module.app.app_context():
        app_module.maybe_sync_master_job_file()
        jobs = {job['number']: job['name'] for job in app_module.get_job_list()}
        assert jobs.get('126028') == '213 Hillandale'
        assert '111111' not in jobs
        assert snapshot_timesheets() == original_rows

        os.utime(master_path, (2_100_000_000, 2_100_000_000))
        app_module._last_master_check = 0
        app_module.maybe_sync_master_job_file()
        jobs = {job['number']: job['name'] for job in app_module.get_job_list()}
        assert jobs.get('111111') == 'Old Master Job'
        assert snapshot_timesheets() == original_rows
        # Put the uploaded list back so the restarted process sees that catalog.
        app_module.save_job_catalog(
            [
                {'number': '126028', 'name': '213 Hillandale'},
                {'number': 'H1000', 'name': 'Overhead'},
                {'number': '122002', 'name': 'Villages of Granite'},
            ],
            source='upload',
        )

    # .xlsx upload path, including a title row above the real headers.
    with open(xlsx_path, 'rb') as handle:
        xlsx_upload = client.post(
            '/admin/upload-jobs',
            data={'file': (handle, 'Job List.xlsx')},
            content_type='multipart/form-data',
        )
    assert xlsx_upload.status_code == 200, xlsx_upload.get_json()
    with app_module.app.app_context():
        jobs = {job['number']: job['name'] for job in app_module.get_job_list()}
        assert jobs == {'126028': '213 Hillandale'}
        assert snapshot_timesheets() == original_rows
        app_module.save_job_catalog(
            [
                {'number': '126028', 'name': '213 Hillandale'},
                {'number': 'H1000', 'name': 'Overhead'},
            ],
            source='upload',
        )

    script = r'''
import os, sys
os.environ["DATABASE_URL"] = sys.argv[1]
os.environ["MASTER_JOB_FILE"] = sys.argv[2]
os.environ["ADMIN_PASSWORD"] = "admin"
import app
with app.app.app_context():
    jobs = {j["number"]: j["name"] for j in app.get_job_list()}
    assert jobs.get("126028") == "213 Hillandale", jobs
    row = app.db.session.get(app.Timesheet, int(sys.argv[3]))
    assert row is not None
    assert row.employee_name == "Brian Clark"
    assert row.job_number == "122001"
    assert row.job_name == "Forest Lake"
    assert row.hours == 4.5
    assert row.description == "• Existing work that must stay"
    assert row.is_billed is True
    assert app.Timesheet.query.count() == 1
    print("persisted", len(jobs))
'''
    subprocess.check_call([
        sys.executable, '-c', script,
        os.environ['DATABASE_URL'],
        str(ROOT / 'missing-master-job-list.xls'),
        str(entry_id),
    ])
    print('job catalog tests passed')


if __name__ == '__main__':
    main()
